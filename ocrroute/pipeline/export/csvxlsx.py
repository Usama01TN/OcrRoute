# coding=utf-8
"""
None
"""
from __future__ import absolute_import, division, print_function

import csv
import io

_HEAD = ['page', 'line', 'word', 'text', 'left', 'top', 'width', 'height', 'confidence']


def _rows(result):
    rows = []
    for i, ln in enumerate(result.get('TextOverlay', {}).get('Lines', []), 1):
        for j, wd in enumerate(ln.get('Words', []), 1):
            rows.append(
                [
                    ln.get('Page', 1),
                    i,
                    j,
                    wd['WordText'],
                    round(wd['Left'], 1),
                    round(wd['Top'], 1),
                    round(wd['Width'], 1),
                    round(wd['Height'], 1),
                    wd.get('Confidence', ''),
                ]
            )
    return rows


def writeCsv(result, meta):
    buf = io.StringIO()
    wr = csv.writer(buf)
    wr.writerow(_HEAD)
    wr.writerows(_rows(result))
    return buf.getvalue().encode('utf-8-sig')


def writeXlsx(result, meta):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = 'Words'
    ws.append(_HEAD)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in _rows(result):
        ws.append(r)
    ws2 = wb.create_sheet('Lines')
    ws2.append(['page', 'line', 'text'])
    for i, ln in enumerate(result.get('TextOverlay', {}).get('Lines', []), 1):
        ws2.append([ln.get('Page', 1), i, ln.get('LineText', '')])
    ws.freeze_panes = 'A2'
    ws2.freeze_panes = 'A2'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
